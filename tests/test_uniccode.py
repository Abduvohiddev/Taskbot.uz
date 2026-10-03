"""@uniccodebot testlari: sqlite (aiosqlite) ustida, Postgres shart emas.

Ishga tushirish: pip install aiosqlite pytest && pytest tests/test_uniccode.py
"""
import asyncio
from datetime import datetime
from zoneinfo import ZoneInfo

import openpyxl
import pytest

pytest.importorskip("aiosqlite")

from uniccode import generator, importer, parsers  # noqa: E402
from uniccode.db import init_db, make_engine, make_sessionmaker  # noqa: E402
from uniccode.excel_out import build_excel  # noqa: E402
from uniccode.generator import Item  # noqa: E402

TZ = ZoneInfo("Asia/Tashkent")
NOW = datetime(2026, 10, 3, 12, 0, tzinfo=TZ)


def run(coro):
    return asyncio.run(coro)


async def _sm(tmp_path):
    engine = make_engine(f"sqlite+aiosqlite:///{tmp_path / 'uc.db'}")
    await init_db(engine)
    return engine, make_sessionmaker(engine)


def test_format_and_normalize():
    assert generator.normalize_article("x-844088") == "844088"
    assert generator.normalize_article(844088.0) == "844088"
    assert generator.date_part(datetime(2026, 9, 25)) == "2509"
    assert generator.format_code("844088", "CA", "2509", 10002897) == "844088CA250910002897"


def test_generate_sequential_and_new_article(tmp_path):
    async def go():
        engine, sm = await _sm(tmp_path)
        async with sm() as s:
            r1 = await generator.generate(s, [Item("844088", 3), Item("844088", 2, address="TM_6B1")], now=NOW)
        async with sm() as s:
            r2 = await generator.generate(s, [Item("844088", 1)], now=NOW)
        await engine.dispose()
        return r1, r2
    r1, r2 = run(go())
    assert r1.items[0].codes == ["844088CA031010000000", "844088CA031010000001", "844088CA031010000002"]
    assert r1.items[1].codes == ["844088CA031010000003", "844088CA031010000004"]
    assert r2.items[0].codes == ["844088CA031010000005"]


def test_rejects_bad_article(tmp_path):
    async def go():
        engine, sm = await _sm(tmp_path)
        async with sm() as s:
            with pytest.raises(ValueError):
                await generator.generate(s, [Item("84408", 1)], now=NOW)
        await engine.dispose()
    run(go())


def test_concurrent_requests_never_collide(tmp_path):
    async def go():
        engine, sm = await _sm(tmp_path)

        async def one():
            async with sm() as s:
                return await generator.generate(s, [Item("814086", 7), Item("825250", 3)], now=NOW)
        results = await asyncio.gather(*[one() for _ in range(10)])
        await engine.dispose()
        return [c for r in results for ir in r.items for c in ir.codes]
    codes = run(go())
    assert len(codes) == 100 and len(set(codes)) == 100


def _legacy_xlsx(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Датабаза"
    ws.append(["Дата"])
    ws.append(["Дата", "Артикул", "Уникальный код", "Уникальный код"])
    ws.append([datetime(2025, 8, 7, 14, 40), 853205, 10000000, "853205CA082510000000"])
    ws.append([datetime(2025, 8, 7, 14, 40), 853205, 10000001, "853205CA082510000001"])
    ws.append([datetime(2025, 8, 7, 14, 40), 853205, 10000001, "853205CA082510000001"])  # takror
    ws.append([datetime(2026, 1, 3, 9, 55), 899209, 10000004, "899209CA03012610000004"])  # 6 xonali sana
    ws.append([datetime(2026, 9, 18, 22, 19), 844088, 10002896, "844088CA180910002896", "TM_7B1"])
    ws.append([datetime(2025, 8, 9, 12, 15), "x-899071", 10000000, "buzuq"])
    uk = wb.create_sheet("Уник код", 0)  # yasalgan, lekin hali saqlanmagan kod
    uk.append(["CA", "0310", None, 844088, 10002976, "844088CA031010002976", None, 20])
    wb.save(path)


def test_import_then_continue(tmp_path):
    src = tmp_path / "baza.xlsx"
    _legacy_xlsx(src)

    async def go():
        engine, sm = await _sm(tmp_path)
        rows, mx, st = importer.read_legacy(str(src))
        async with sm() as s:
            st = await importer.import_rows(s, rows, mx, st)
        rows2, mx2, st2 = importer.read_legacy(str(src))  # qayta import xavfsiz
        async with sm() as s:
            st2 = await importer.import_rows(s, rows2, mx2, st2)
        async with sm() as s:
            res = await generator.generate(s, [Item("844088", 2), Item("899209", 1), Item("899071", 1)], now=NOW)
        await engine.dispose()
        return st, st2, res
    st, st2, res = run(go())
    assert st.inserted == 5 and st.duplicates_in_file == 1 and st.irregular == 1 and st.invalid == 0
    assert st2.inserted == 0
    assert st.pending_reserved == 1
    assert res.items[0].codes == ["844088CA031010002977", "844088CA031010002978"]
    assert res.items[1].codes == ["899209CA031010000005"]
    assert res.items[2].codes == ["899071CA031010000001"]  # buzuq kod qatoridagi raqam ham hisobga olindi


def _stock_rows():
    hdr = ["Артикул", None, "Номенклатура", None, None, None, "Адрес", "Начальный остаток", "Приход",
           "Расход", "Конечный остаток", "К отбору", "К размещению"]
    def line(art, name, addr, end):
        return [art, None, name, None, None, None, addr, end, None, None, end, None, None]
    return [
        [None] * 13,
        ["Движения товаров на адресных складах"] + [None] * 12,
        ["Параметры:", None, None, "Период: 01/10/2026 - 01/10/2026"] + [None] * 9,
        ["Адрес"] + [None] * 6 + ["Количество упаковок"] + [None] * 5,
        hdr,
        ["TM_9A1"] + [None] * 12,
        line(804066, "X-Co/244/Rich-Ultra/Qaymoq/AL-01/KJ-01/#00", "TM_9A1", 2),
        line(899093, "C-Universal Nakidka Palma/ToqChoko", "TM_9A1", 3),
        line(390181, "FocusEva Cobalt 5D Vanna", "TM_1A1", None),
        line(814220, "C-Co-220/Myunxen/IB8877-01/KJ-01/032", "S_TM-Transit", 5),
        line(450000, "C-Nezamerzayka-20-3L", "S_TM-Transit", 7),
        ["Итого"] + [None] * 12,
    ]


def test_stock_report_selection():
    rep = parsers.parse_stock_report(_stock_rows())
    assert rep and len(rep.lines) == 5 and "01/10/2026" in rep.period
    chosen, skipped = parsers.select_lines(rep, "9,1", only_chexol=False)
    assert [(l.article, l.qty) for l in chosen] == [("804066", 2), ("899093", 3)]
    assert [l.article for l in skipped] == ["390181"]
    chosen, _ = parsers.select_lines(rep, "TM_9A1 S_TM-Transit", only_chexol=True)
    assert [l.article for l in chosen] == ["804066", "814220"]  # nakidka va nezamerzayka chiqmaydi
    assert rep.rows_summary() == {"9": 5, "S_TM-Transit": 12}


def test_list_file_with_series_and_errors():
    rows = [["Артикул", "Кол-во", "Seria", "Наименование"],
            [844088, 3, "T535", "CY-Co-152"],
            ["x-814086", "2", "K1807", None],
            [12345, 1, None, None],
            [844205, 0, None, None],
            ["Итого", 5, None, None]]
    lst = parsers.parse_list(rows)
    assert [(i.article, i.qty, i.series) for i in lst.items] == [("844088", 3, "T535"), ("814086", 2, "K1807")]
    assert lst.items[0].note == "T535"
    assert len(lst.errors) == 2
    headerless = parsers.parse_list([[844088, 2, "P1"], [814086, 1, None]])
    assert [(i.article, i.qty) for i in headerless.items] == [("844088", 2), ("814086", 1)]


def test_excel_output(tmp_path):
    async def go():
        engine, sm = await _sm(tmp_path)
        async with sm() as s:
            r = await generator.generate(s, [Item("844088", 2, address="TM_6B1", series="T1")], now=NOW)
        await engine.dispose()
        return r
    res = run(go())
    p = tmp_path / "o.xlsx"
    p.write_bytes(build_excel(res, title="test"))
    wb = openpyxl.load_workbook(p)
    ws = wb["Kodlar"]
    assert [c.value for c in ws[2]][1:5] == [844088, 10000000, "844088CA031010000000", "TM_6B1"]
    assert wb["Xulosa"].cell(3, 5).value == 2


def test_sheet_sync_uploads_once_and_retries(tmp_path):
    from uniccode.config import UnicSettings
    from uniccode.sheets_sync import SheetSync

    async def go():
        engine, sm = await _sm(tmp_path)
        async with sm() as s:
            await generator.generate(s, [Item("844088", 3, address="TM_6B1")], now=NOW)
        sync = SheetSync(sm, UnicSettings(GSHEET_BATCH=2))
        uploaded, fail = [], {"on": True}

        def fake_append(values):
            if fail["on"]:
                raise ConnectionError("google ishlamayapti")
            uploaded.extend(values)
        sync._append = fake_append
        with pytest.raises(ConnectionError):
            await sync.sync_once()            # xato bo'lsa kodlar kutib turadi
        assert await sync.pending() == 3
        fail["on"] = False
        assert await sync.sync_once() == 2    # GSHEET_BATCH=2
        assert await sync.sync_once() == 1
        assert await sync.sync_once() == 0    # takror yuklanmaydi
        await engine.dispose()
        return uploaded
    up = run(go())
    assert [v[3] for v in up] == ["844088CA031010000000", "844088CA031010000001", "844088CA031010000002"]
    assert up[0][0] == "03/10/2026 12:00:00" and up[0][4] == "TM_6B1"


def test_catalog_and_series_files(tmp_path):
    import xlrd
    from uniccode import catalog
    from uniccode.excel_out import build_rfid_xls, build_seria

    src = tmp_path / "kat.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Артикул ", "Номенклатура", "Номенклатура.Avto Marka"])
    ws.append([814292, "C-Cobalt-229/Matrix", "COB"])
    ws.append(["x-844088", "CY-Co-152/Alfa", None])
    ws.append([12345, "buzuq", None])
    ws.append([899000, 0, None])  # nomi yo'q - o'tkaziladi
    ws2 = wb.create_sheet("DATABASE")  # sarlavhasiz varaq
    ws2.append([846088, "CY-L3-152/Alfa"])
    ws2.append([814292, "boshqa nom"])   # birinchi varaq ustun turadi
    wb.save(src)
    found = catalog.parse_catalog(str(src))
    assert found == {"814292": ("C-Cobalt-229/Matrix", "COB"), "844088": ("CY-Co-152/Alfa", None),
                     "846088": ("CY-L3-152/Alfa", None)}

    async def go():
        engine, sm = await _sm(tmp_path)
        async with sm() as s:
            assert await catalog.save_catalog(s, found) == 3
        async with sm() as s:
            prods = await catalog.get_products(s, ["814292", "999999"])
        async with sm() as s:
            res = await generator.generate(s, [
                Item("814292", 2, name=prods["814292"].name, series="F-1795", machine="COB", note="F-1795"),
                Item("844088", 1, name="CY-Co-152/Alfa", series="F-1795", machine="NEX", note="F-1795"),
            ], now=NOW)
        await engine.dispose()
        return prods, res
    prods, res = run(go())
    assert set(prods) == {"814292"} and prods["814292"].marka == "COB"

    p = tmp_path / "s.xlsx"
    p.write_bytes(build_seria(res))
    ws = openpyxl.load_workbook(p).active
    assert [c.value for c in ws[1]][:3] == ["artikul", "nomeklatura ", "konveyr"]
    assert [c.value for c in ws[2]][:7] == [814292, "C-Cobalt-229/Matrix", "F-1795", "COB", "F-1795-COB",
                                            10000000, "814292CA031010000000"]
    assert ws.cell(4, 5).value == "F-1795-NEX" and ws.max_row == 4
    assert ws.cell(1, 9).value is None and [ws.cell(r, 9).value for r in (2, 3, 4)] == [2, 2, 2]

    x = tmp_path / "r.xls"
    x.write_bytes(build_rfid_xls(res))
    sh = xlrd.open_workbook(str(x)).sheet_by_name("STA")
    assert sh.row_values(0)[:4] == ["Artikul", "EPC", "Nomi", "ummumiy  nomi "]
    assert sh.row_values(1)[:4] == [814292.0, "814292CA031010000000", "C-Cobalt-229/Matrix", "F-1795-COB"]
    assert sh.nrows == 4


def test_onec_fetch_paging_retry_and_auth():
    import base64
    from aiohttp import web
    from uniccode import onec

    rows = [{"Ref_Key": f"k{i:04d}", "Артикул": str(810000 + i), "Description": f"qisqa {i}",
             "НаименованиеПолное": f"To'liq nom {i}" if i % 2 else ""} for i in range(25)]
    rows += [{"Ref_Key": "z1", "Артикул": "13072015", "Description": "x", "НаименованиеПолное": "x"},
             {"Ref_Key": "z2", "Артикул": "x-899071", "Description": "Alkantaro", "НаименованиеПолное": ""}]
    seen = {"fail_once": True, "params": []}

    async def handler(request):
        auth = request.headers.get("Authorization", "")
        if auth != "Basic " + base64.b64encode(b"user:pass").decode():
            return web.Response(status=401)
        q = request.rel_url.query
        seen["params"].append(dict(q))
        if seen["fail_once"]:
            seen["fail_once"] = False
            return web.Response(status=500, text="Информационная база не обнаружена")
        top, skip = int(q["$top"]), int(q["$skip"])
        return web.json_response({"value": rows[skip:skip + top]})

    async def go():
        app = web.Application()
        app.router.add_get("/ERP25/odata/standard.odata/Catalog_Номенклатура", handler)
        runner = web.AppRunner(app)
        await runner.setup()
        site = web.TCPSite(runner, "127.0.0.1", 0)
        await site.start()
        port = site._server.sockets[0].getsockname()[1]
        base = f"http://127.0.0.1:{port}/ERP25/odata/standard.odata/"
        try:
            got = await onec.fetch_catalog(base, "user", "pass", page=10, delay=0.01)
            with pytest.raises(onec.OneCError):
                await onec.fetch_catalog(base, "user", "xato", page=10, delay=0.01)
        finally:
            await runner.cleanup()
        return got
    got = run(go())
    assert len(got) == 26                       # 25 + x-899071, 8 xonali artikul tashlandi
    assert got["810001"] == ("To'liq nom 1", None) and got["810002"] == ("qisqa 2", None)
    assert got["899071"] == ("Alkantaro", None)
    p = seen["params"][-1]
    assert p["$orderby"] == "Ref_Key" and p["$filter"] == "IsFolder eq false and DeletionMark eq false"
    assert [x["$skip"] for x in seen["params"][1:4]] == ["0", "10", "20"]   # 500 dan keyin qayta urinish
