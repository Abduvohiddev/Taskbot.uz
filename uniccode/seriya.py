"""
1C: har bir vid nomenklatura uchun seriyalar sozlamasi (qaysi skladda qaysi "Политика учета серий").

Manbalar (faqat o'qish, OData):
- Catalog_ВидыНоменклатуры      - vidlar, "Серии" belgisi, umumiy politika va skladlar bo'yicha individual politikalar
                                   (jadval qism: Склад + Политика). Maydon nomlari konfiguratsiyaga qarab
                                   biroz farq qilishi mumkin, shuning uchun ular nomidan topiladi.
- Catalog_ПолитикиУчетаСерий    - politika nomlari ("Управление остатками серий 1" va h.k.)
- Catalog_Склады                - sklad nomlari
- Catalog_Номенклатура          - artikullar va ularning vidi
"""
import re
from collections import defaultdict
from io import BytesIO
from typing import Dict, List, Optional
from urllib.parse import quote

import aiohttp
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from uniccode.generator import normalize_article
from uniccode.onec import NotPublished, OneCError, _get_page

SERIYA1_RE = re.compile(r"сери[йи]\s*1(?!\d)", re.I)  # "Управление остатками серий 1"
EMPTY_KEY = "00000000-0000-0000-0000-000000000000"
NO_SERIES = "Серии не используются"

# Catalog_ПолитикиУчетаСерий OData'ga ochilmagan bo'lsa, politika nomlari 1C dagi ma'lum namunadan
# aniqlanadi (Cardinar Chexol vidi kartochkasi): (vid, sklad nomidagi so'z yoki None = umumiy politika, nom).
ANCHORS = [
    ("Cardinar Chexol", None, NO_SERIES),
    ("Cardinar Chexol", "Sergeli Chexol TM", "Управление остатками серий 1"),
    ("Cardinar Chexol", "Ochrilgan", "_Управление по FEFO остатками серий"),
]


async def fetch_all(session: aiohttp.ClientSession, base_url: str, entity: str, params: dict,
                    page: int = 500, retries: int = 4, delay: float = 5.0) -> List[dict]:
    url = base_url.rstrip("/") + "/" + quote(entity)
    out: List[dict] = []
    skip = 0
    while True:
        p = {"$format": "json", **params, "$orderby": "Ref_Key", "$top": str(page), "$skip": str(skip)}
        rows = await _get_page(session, url, p, retries, delay)
        out.extend(rows)
        if len(rows) < page:
            return out
        skip += page


def _key(v) -> Optional[str]:
    return v if isinstance(v, str) and v and v != EMPTY_KEY else None


def _find_field(row: dict, *words, suffix="", typ=None) -> Optional[str]:
    """Nomida hamma so'zlar bor birinchi maydon (masalan 'Политика', 'Сери' + '_Key')."""
    for k, v in row.items():
        kl = k.lower()
        if all(w.lower() in kl for w in words) and k.endswith(suffix) and (typ is None or isinstance(v, typ)):
            return k
    return None


def parse_vid(row: dict) -> dict:
    """Bitta vid yozuvidan seriya sozlamalarini ajratib oladi."""
    uses = row.get("ИспользоватьСерии")
    if uses is None:
        f = _find_field(row, "Использ", "Сери", typ=bool)
        uses = row.get(f) if f else None
    pol_f = _find_field(row, "Политик", "Сери", suffix="_Key")
    indiv_f = _find_field(row, "Индивидуальн", typ=bool)
    per_sklad = []
    for k, v in row.items():
        if not (isinstance(v, list) and v and isinstance(v[0], dict)):
            continue
        sk_f = next((f for f in v[0] if "склад" in f.lower() and f.endswith("_Key")), None)
        if not sk_f:
            continue
        pf = next((f for f in v[0] if "полит" in f.lower() and f.endswith("_Key")), None) or \
            next((f for f in v[0] if f.endswith("_Key") and f != sk_f), None)
        for line in v:
            per_sklad.append((_key(line.get(sk_f)), _key(line.get(pf)) if pf else None))
    return dict(
        key=row.get("Ref_Key"), name=(row.get("Description") or "").strip(), parent=_key(row.get("Parent_Key")),
        folder=bool(row.get("IsFolder")), deleted=bool(row.get("DeletionMark")),
        uses=uses, policy=_key(row.get(pol_f)) if pol_f else None,
        individual=row.get(indiv_f) if indiv_f else None, per_sklad=per_sklad,
    )


async def fetch_series_data(base_url: str, user: str, password: str, page: int = 500,
                            delay: float = 5.0, timeout: float = 300.0) -> dict:
    auth = aiohttp.BasicAuth(user, password, encoding="utf-8") if user else None
    async with aiohttp.ClientSession(auth=auth, timeout=aiohttp.ClientTimeout(total=timeout)) as s:
        vids = await fetch_all(s, base_url, "Catalog_ВидыНоменклатуры", {}, page, delay=delay)
        if not vids:
            raise OneCError("1C bo'sh ro'yxat qaytardi (Catalog_ВидыНоменклатуры).")
        missing = []

        async def optional(entity):
            try:
                return await fetch_all(s, base_url, entity, {"$select": "Ref_Key,Description"}, page, delay=delay)
            except NotPublished:
                missing.append(entity)
                return None
        pols = await optional("Catalog_ПолитикиУчетаСерий")
        sklads = await optional("Catalog_Склады")
        noms = await fetch_all(s, base_url, "Catalog_Номенклатура", {
            "$filter": "IsFolder eq false and DeletionMark eq false",
            "$select": "Ref_Key,Артикул,Description,НаименованиеПолное,ВидНоменклатуры_Key",
        }, 1000, delay=delay)
    return dict(vids=vids, pols=pols, sklads=sklads, noms=noms, missing=missing)


def policy_names(data: dict, parsed: List[dict], sk_name: Dict[str, str]):
    """{politika_key: nom} va qanday aniqlangani haqida izohlar."""
    if data.get("pols") is not None:
        return {p["Ref_Key"]: (p.get("Description") or "").strip() for p in data["pols"]}, []
    names: Dict[str, str] = {}
    notes = ["Catalog_ПолитикиУчетаСерий OData'ga ochilmagan: politika nomlari Cardinar Chexol namunasidan aniqlandi."]
    for vid, sklad_word, label in ANCHORS:
        v = next((p for p in parsed if p["name"].lower() == vid.lower() and not p["folder"]), None)
        if not v:
            continue
        if sklad_word is None:
            key = v["policy"]
        else:
            key = next((pk for sk, pk in v["per_sklad"]
                        if sklad_word.lower() in sk_name.get(sk, "").lower()), None)
        if key and key not in names:
            names[key] = label
            notes.append(f"{label} = {vid} / {sklad_word or 'umumiy'} dagi politika ({key[:8]}…)")
    seen = sorted({k for p in parsed for k in [p["policy"]] + [pk for _, pk in p["per_sklad"]] if k and k not in names})
    for i, k in enumerate(seen, 1):
        names[k] = f"Nomi noma'lum politika №{i} ({k[:8]}…)"
    return names, notes


HF = Font(name="Arial", bold=True)
BF = Font(name="Arial")
HEAD = PatternFill("solid", fgColor="FFD966")
GREEN = PatternFill("solid", fgColor="C6EFCE")
RED = PatternFill("solid", fgColor="FFC7CE")
GRAY = PatternFill("solid", fgColor="EDEDED")


def _table(ws, header, rows, widths):
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill = HF, HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = BF
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "C2" if len(header) > 6 else "A2"
    ws.auto_filter.ref = ws.dimensions


def build_report(data: dict) -> bytes:
    """Excel: Vidlar (sklad bo'yicha jadval), Batafsil, Artikullar, Diagnostika."""
    sk_name = {s["Ref_Key"]: (s.get("Description") or "").strip() for s in data.get("sklads") or []}
    parsed = [parse_vid(v) for v in data["vids"]]
    pol_name, notes = policy_names(data, parsed, sk_name)
    by_key = {p["key"]: p for p in parsed}
    vids = sorted((p for p in parsed if not p["folder"] and not p["deleted"]),
                  key=lambda p: ((by_key.get(p["parent"]) or {}).get("name", ""), p["name"].lower()))

    arts: Dict[str, list] = defaultdict(list)
    for n in data["noms"]:
        a = normalize_article(n.get("Артикул"))
        arts[n.get("ВидНоменклатуры_Key")].append(
            (a, (n.get("НаименованиеПолное") or n.get("Description") or "").strip()))

    def pname(k):
        return pol_name.get(k, k or "")

    def is1(name):
        return bool(name and SERIYA1_RE.search(name))

    sklad_cols = sorted({sk for v in vids for sk, _ in v["per_sklad"] if sk}, key=lambda k: sk_name.get(k, k))
    status = {}
    rows, detail = [], []
    for v in vids:
        group = (by_key.get(v["parent"]) or {}).get("name", "")
        default = pname(v["policy"]) if v["policy"] else (NO_SERIES if v["uses"] is False else "")
        per = {sk: pname(pk) for sk, pk in v["per_sklad"] if sk}
        yes = [sk_name.get(sk, sk) for sk, pn in per.items() if is1(pn)]
        no = [sk_name.get(sk, sk) for sk, pn in per.items() if not is1(pn)]
        if is1(default):
            yes.insert(0, "(umumiy - hamma skladlar)")
        st = "Ha" if yes else "Yo'q"
        status[v["key"]] = (st, ", ".join(yes))
        uses = {True: "Ha", False: "Yo'q"}.get(v["uses"], "?")
        rows.append([group, v["name"], uses, default, len(arts.get(v["key"], [])), st, ", ".join(yes), ", ".join(no)]
                    + [per.get(sk, "") for sk in sklad_cols])
        for sk, pn in per.items():
            detail.append([group, v["name"], sk_name.get(sk, sk), pn, "Ha" if is1(pn) else "Yo'q"])
        if not per:
            detail.append([group, v["name"], "(individual sozlama yo'q)", default, "Ha" if is1(default) else "Yo'q"])

    wb = Workbook()
    ws = wb.active
    ws.title = "Vidlar"
    head = ["Группа видов", "Вид номенклатуры", "Серии", "Umumiy politika", "Artikullar",
            "Seriya 1 yoqilgan", "Seriya 1 yoqilgan skladlar", "Boshqa politika (seriya 1 emas)"] \
        + [sk_name.get(sk, sk) for sk in sklad_cols]
    _table(ws, head, rows, [16, 30, 8, 26, 10, 10, 40, 40] + [22] * len(sklad_cols))
    for row in ws.iter_rows(min_row=2):
        row[5].fill = GREEN if row[5].value == "Ha" else RED
        for c in row[8:]:
            if c.value:
                c.fill = GREEN if is1(c.value) else GRAY

    ws2 = wb.create_sheet("Batafsil")
    _table(ws2, ["Группа видов", "Вид номенклатуры", "Склад", "Политика учета серий", "Seriya 1"],
           detail, [16, 30, 34, 40, 10])
    for row in ws2.iter_rows(min_row=2):
        row[4].fill = GREEN if row[4].value == "Ha" else RED

    ws3 = wb.create_sheet("Artikullar")
    arows = []
    vid_name = {v["key"]: v for v in vids}
    for k, items in arts.items():
        v = vid_name.get(k)
        name = v["name"] if v else "(vid ko'rsatilmagan)"
        group = (by_key.get(v["parent"]) or {}).get("name", "") if v else ""
        st, where = status.get(k, ("Yo'q", ""))
        for a, nm in items:
            arows.append([group, name, int(a) if a.isdigit() else a, nm, st, where])
    arows.sort(key=lambda r: (r[0], r[1].lower(), str(r[2])))
    _table(ws3, ["Группа видов", "Вид номенклатуры", "Артикул", "Наименование", "Seriya 1 yoqilgan", "Qaysi skladlarda"],
           arows, [16, 30, 10, 60, 10, 40])
    for row in ws3.iter_rows(min_row=2):
        row[4].fill = GREEN if row[4].value == "Ha" else RED

    ws4 = wb.create_sheet("Diagnostika")
    sample = next((v for v in data["vids"] if not v.get("IsFolder")), data["vids"][0])
    for line in [
        f"Vidlar: {len(vids)} (papkalar bilan jami {len(parsed)})",
        f"Politikalar: {len(pol_name)}; skladlar: {len(sk_name)}; nomenklatura: {len(data['noms'])}",
        "OData'ga ochilmagan: " + (", ".join(data.get("missing") or []) or "yo'q"),
        *notes,
        f"Skladlar bo'yicha individual sozlamasi bor vidlar: {sum(1 for v in vids if v['per_sklad'])}",
        "Vid yozuvi maydonlari: " + ", ".join(
            f"{k}[{len(v)}]" if isinstance(v, list) else k for k, v in sample.items()),
        "Politikalar: " + "; ".join(sorted(pol_name.values())),
    ]:
        ws4.append([line])
    ws4.column_dimensions["A"].width = 160

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def summary(data: dict) -> str:
    parsed = [parse_vid(v) for v in data["vids"]]
    sk_name = {s["Ref_Key"]: (s.get("Description") or "").strip() for s in data.get("sklads") or []}
    pol_name, notes = policy_names(data, parsed, sk_name)
    vids = [v for v in parsed if not v["folder"] and not v["deleted"]]
    yes = sum(1 for v in vids
              if SERIYA1_RE.search(pol_name.get(v["policy"], "") or "")
              or any(SERIYA1_RE.search(pol_name.get(pk, "") or "") for _, pk in v["per_sklad"]))
    txt = f"Vidlar: {len(vids)}. Seriya 1 yoqilgan: {yes}, yoqilmagan: {len(vids) - yes}."
    if data.get("missing"):
        txt += ("\n⚠️ 1C OData'ga ochilmagan: " + ", ".join(data["missing"]) +
                ". Nomlar namunadan aniqlandi, 'Diagnostika' varag'iga qarang.")
        if "Catalog_Склады" in data["missing"]:
            txt += " Sklad nomlari o'rniga kodlar chiqadi."
    return txt
