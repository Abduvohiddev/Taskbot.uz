"""
Eski bazani (Google Sheets "Датабаза" varag'i, xlsx qilib yuklab olingan) SQL ga ko'chirish.

- Har bir kod uc_codes ga yoziladi, bazada bor kodlar o'tkazib yuboriladi (takror yozilmaydi).
- Har bir artikul uchun hisoblagich max(hozirgi, fayldagi eng katta raqam) ga ko'tariladi.
  Faylni qayta-qayta import qilish xavfsiz.
"""
import re
from dataclasses import dataclass, field
from datetime import datetime
from typing import Dict, List, Optional
from zoneinfo import ZoneInfo

import openpyxl
from sqlalchemy import func
from sqlalchemy.dialects import postgresql, sqlite
from sqlalchemy.ext.asyncio import AsyncSession

from uniccode.generator import normalize_article
from uniccode.models import UCBatch, UCCode, UCCounter

# Eski bazada sana qismi ba'zan 6 xonali (030126) - shuni ham qabul qilamiz.
CODE_RE = re.compile(r"^(\d{6})([A-Z]{2})(\d{4}|\d{6})(\d{8})$")


@dataclass
class ImportStats:
    rows_read: int = 0
    valid: int = 0
    invalid: int = 0
    irregular: int = 0
    duplicates_in_file: int = 0
    inserted: int = 0
    articles: int = 0
    invalid_samples: List[str] = field(default_factory=list)


def _find_db_sheet(wb):
    for name in ("Датабаза", "Databaza", "Database"):
        if name in wb.sheetnames:
            return wb[name]
    for ws in wb.worksheets:  # "Уникальный код" sarlavhasi bor birinchi varaq
        for row in ws.iter_rows(min_row=1, max_row=5, values_only=True):
            if row and any(isinstance(v, str) and "Уникальный код" in v for v in row):
                return ws
    raise ValueError("Faylda 'Датабаза' varag'i topilmadi")


def _to_dt(v, tz) -> datetime:
    if isinstance(v, datetime):
        return v if v.tzinfo else v.replace(tzinfo=tz)
    if isinstance(v, str):
        for fmt in ("%Y-%m-%dT%H:%M:%S", "%Y-%m-%d %H:%M:%S", "%d/%m/%Y %H:%M:%S", "%d.%m.%Y %H:%M:%S", "%d/%m/%Y", "%d.%m.%Y"):
            try:
                return datetime.strptime(v.strip()[:19], fmt).replace(tzinfo=tz)
            except ValueError:
                continue
    return datetime(2025, 1, 1, tzinfo=tz)


def read_legacy(path: str, tz_name: str = "Asia/Tashkent"):
    """Faylni o'qiydi: (kodlar ro'yxati, artikul -> eng katta raqam, statistika)."""
    tz = ZoneInfo(tz_name)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = _find_db_sheet(wb)
    stats = ImportStats()
    seen = set()
    rows: List[dict] = []
    max_seq: Dict[str, int] = {}
    for r in ws.iter_rows(min_row=1, values_only=True):
        if not r or len(r) < 4:
            continue
        date_v, art_v, seq_v, code_v = r[0], r[1], r[2], r[3]
        note_v = r[4] if len(r) > 4 else None
        code = str(code_v).strip() if code_v is not None else ""
        if not code or code == "Уникальный код":
            continue
        stats.rows_read += 1
        m = CODE_RE.match(code)
        # Artikul/raqam ustunlari ham hisoblagichga ta'sir qiladi (kod buzuq bo'lsa ham).
        art_col = normalize_article(art_v)
        try:
            seq_col = int(float(seq_v)) if seq_v is not None else None
        except (TypeError, ValueError):
            seq_col = None
        if re.match(r"^\d{6}$", art_col) and seq_col:
            max_seq[art_col] = max(max_seq.get(art_col, 0), seq_col)
        if m:
            art, seq = m.group(1), int(m.group(4))
            max_seq[art] = max(max_seq.get(art, 0), seq)
            stats.valid += 1
        elif re.match(r"^\d{6}$", art_col) and seq_col:
            # Nostandart kod (masalan 'x-899063CA...'): kodning o'zi saqlanadi, artikul/raqam ustunlardan olinadi.
            art, seq = art_col, seq_col
            stats.irregular += 1
        else:
            stats.invalid += 1
            if len(stats.invalid_samples) < 10:
                stats.invalid_samples.append(code)
            continue
        if code in seen:
            stats.duplicates_in_file += 1
            continue
        seen.add(code)
        note = str(note_v).strip() if note_v not in (None, "") else None
        rows.append(dict(
            code=code[:32], article=art, seq=seq, created_at=_to_dt(date_v, tz),
            note=(note[:255] if note else None), source="import", synced=True,
        ))
    wb.close()
    stats.articles = len(max_seq)
    return rows, max_seq, stats


async def import_rows(session: AsyncSession, rows: List[dict], max_seq: Dict[str, int],
                      stats: ImportStats, user_id: Optional[int] = None, filename: Optional[str] = None) -> ImportStats:
    pg = session.bind.dialect.name == "postgresql"
    ins = postgresql.insert if pg else sqlite.insert
    batch = UCBatch(user_id=user_id, source="import", filename=filename)
    session.add(batch)
    await session.flush()
    for r in rows:
        r["batch_id"] = batch.id
    inserted = 0
    for i in range(0, len(rows), 2500):
        chunk = rows[i:i + 2500]
        res = await session.execute(
            ins(UCCode).values(chunk).on_conflict_do_nothing(index_elements=[UCCode.code]).returning(UCCode.id)
        )
        inserted += len(res.fetchall())
    greatest = func.greatest if pg else func.max  # sqlite da max(a, b) skalyar funksiya
    items = list(max_seq.items())
    for i in range(0, len(items), 2000):
        vals = [dict(article=a, last_seq=s) for a, s in items[i:i + 2000]]
        stmt = ins(UCCounter).values(vals)
        stmt = stmt.on_conflict_do_update(
            index_elements=[UCCounter.article],
            set_={"last_seq": greatest(UCCounter.last_seq, stmt.excluded.last_seq), "updated_at": func.now()},
        )
        await session.execute(stmt)
    batch.count = inserted
    await session.commit()
    stats.inserted = inserted
    return stats
