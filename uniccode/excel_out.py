"""Natija Excel fayli: Датабаза formatidagi kodlar + xulosa."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from uniccode.generator import BatchResult

HF = Font(name="Arial", bold=True)
BF = Font(name="Arial")
FILL = PatternFill("solid", fgColor="DDEBF7")


def _sheet(ws, header, rows, widths):
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill, c.alignment = HF, FILL, Alignment(horizontal="center")
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = BF
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def build_excel(result: BatchResult, title: str = "") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Kodlar"
    created = result.created_at.replace(tzinfo=None)
    rows = []
    for ir in result.items:
        it = ir.item
        for i, code in enumerate(ir.codes):
            rows.append([created, int(it.article), ir.first_seq + i, code, it.note or it.address or it.series,
                         it.address, it.name, it.series])
    _sheet(ws, ["Дата", "Артикул", "Уникальный код", "Уникальный код", "Izoh", "Адрес", "Наименование", "Seria"],
           rows, [20, 10, 16, 24, 16, 15, 55, 16])
    for row in ws.iter_rows(min_row=2, max_col=1):
        row[0].number_format = "DD/MM/YYYY HH:MM:SS"

    ws2 = wb.create_sheet("Xulosa")
    srows = [[ir.item.address, int(ir.item.article), ir.item.name, ir.item.series, len(ir.codes),
              ir.codes[0] if ir.codes else None, ir.codes[-1] if ir.codes else None] for ir in result.items]
    srows.append(["JAMI", None, None, None, result.total, None, None])
    _sheet(ws2, ["Адрес", "Артикул", "Наименование", "Seria", "Kod soni", "Birinchi kod", "Oxirgi kod"],
           srows, [15, 10, 55, 16, 10, 24, 24])
    ws2.cell(ws2.max_row, 1).font = HF
    ws2.cell(ws2.max_row, 5).font = HF
    if title:
        ws3 = wb.create_sheet("Izoh")
        ws3.column_dimensions["A"].width = 120
        for line in title.split("\n"):
            ws3.append([line])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------- seriya va RFID fayllari
YELLOW = PatternFill("solid", fgColor="FFFF00")
SERIA_HEADER = ["artikul", "nomeklatura ", "konveyr", "mashinasi", "ummumiy nomi ", "unic code", "unic code full", "sana",
                None]  # I ustuni: shablonda sarlavhasiz, har qatorda bir xil son (UNIC_SERIA_I, odatda 2)


def umumiy_nomi(konveyr, mashina) -> str:
    return "-".join(x for x in (konveyr or "", mashina or "") if x)


def build_seria(result: BatchResult, i_value=2) -> bytes:
    """Seriya fayli (seria_1.xlsx shabloni bo'yicha). Sariq ustunlar: artikul, konveyr, mashinasi."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Лист1"
    ws.append(SERIA_HEADER)
    for i, c in enumerate(ws[1]):
        c.font = HF
        if i in (0, 2, 3):
            c.fill = YELLOW
    day = result.created_at.replace(tzinfo=None, hour=0, minute=0, second=0, microsecond=0)
    for ir in result.items:
        it = ir.item
        for i, code in enumerate(ir.codes):
            ws.append([int(it.article), it.name or "", it.series or "", it.machine or "",
                       umumiy_nomi(it.series, it.machine), ir.first_seq + i, code, day, i_value])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = BF
        row[7].number_format = "dd/mm/yyyy"
    for col, w in zip("ABCDEFGHI", [9, 60, 12, 13, 18, 12, 24, 12, 6]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_rfid_xls(result: BatchResult) -> bytes:
    """RFID ilovasi uchun .xls ('STA' varag'i): Artikul | EPC | Nomi | ummumiy nomi + 8 bo'sh ustun."""
    import xlwt
    total = sum(len(ir.codes) for ir in result.items)
    if total > 65000:
        raise ValueError(f"RFID .xls fayliga {total} qator sig'maydi (chegara 65000). Bo'lib yuboring.")
    wb = xlwt.Workbook(encoding="utf-8")
    ws = wb.add_sheet("STA")
    red = xlwt.easyxf("pattern: pattern solid, fore_colour red; font: bold on")
    gray = xlwt.easyxf("pattern: pattern solid, fore_colour gray25; font: bold on")
    header = ["Artikul", "EPC", "Nomi", "ummumiy  nomi "] + [""] * 8
    for j, h in enumerate(header):
        ws.write(0, j, h, red if j < 2 else gray)
    r = 1
    for ir in result.items:
        it = ir.item
        un = umumiy_nomi(it.series, it.machine)
        for code in ir.codes:
            ws.write(r, 0, int(it.article))
            ws.write(r, 1, code)
            ws.write(r, 2, it.name or "")
            ws.write(r, 3, un)
            r += 1
    for j, w in enumerate([10, 24, 60, 18]):
        ws.col(j).width = 256 * w
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
